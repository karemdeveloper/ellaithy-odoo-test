"""Bounded XLSX parsing, deliberately independent of Odoo and the database."""
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
import re
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


HEADERS = {
    'code': ('كود الموظف', 'employee code', 'code'),
    'date': ('التاريخ', 'attendance date', 'date'),
    'check_in': ('الحضور', 'check in', 'check-in', 'check_in'),
    'check_out': ('الانصراف', 'انصراف', 'check out', 'check-out', 'check_out'),
    'shift': ('الشيفت', 'shift'),
}
LEAVE = {'إجازة', 'اجازة', 'leave', 'vacation', 'غياب', 'absent'}


def open_excel(data, max_mb=10):
    if not data or len(data) > max_mb * 1024 * 1024:
        raise ValueError('Empty file or configured file size limit exceeded.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > max_mb * 20 * 1024 * 1024:
                raise ValueError('Expanded workbook exceeds the safety limit.')
        return load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError, TypeError, SyntaxError, ParseError, InvalidFileException) as exc:
        raise ValueError('Invalid XLSX workbook.') from exc


def employee_code(cell):
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, bool):
        raise ValueError('Employee code must be text or an integer.')
    if isinstance(value, (int, float)):
        if int(value) != value or abs(value) >= 10**15:
            raise ValueError('Store long or non-integer employee codes as Excel text.')
        result = str(int(value))
        if re.fullmatch(r'0+', cell.number_format or ''):
            result = result.zfill(len(cell.number_format))
        return result
    return str(value).strip()


def read_rows(book, sheet_number=1, max_rows=10000):
    try:
        return _read_rows(book, sheet_number, max_rows)
    except (BadZipFile, SyntaxError, ParseError) as exc:
        raise ValueError('Invalid worksheet XML.') from exc


def _read_rows(book, sheet_number=1, max_rows=10000):
    if not 1 <= sheet_number <= len(book.worksheets):
        raise ValueError('Select a valid worksheet number.')
    sheet = book.worksheets[sheet_number - 1]
    if sheet.max_column and sheet.max_column > 100:
        raise ValueError('Maximum 100 columns supported.')
    if sheet.max_row and sheet.max_row > max_rows + 1:
        raise ValueError('Configured worksheet row limit exceeded.')
    iterator = sheet.iter_rows()
    header = next(iterator, ())
    names = [str(c.value or '').strip().lower() for c in header]
    mapping = {}
    for key, aliases in HEADERS.items():
        matches = [i for i, name in enumerate(names) if name in aliases]
        if len(matches) > 1:
            raise ValueError('Ambiguous header: %s' % key)
        if matches:
            mapping[key] = matches[0]
        elif key != 'shift':
            raise ValueError('Missing required column: %s' % key)
    rows = []
    for number, cells in enumerate(iterator, 2):
        if number > max_rows + 1:
            raise ValueError('Configured worksheet row limit exceeded.')
        if all(c.value is None or str(c.value).strip() == '' for c in cells):
            continue
        values = {key: cells[index].value if index < len(cells) else None for key, index in mapping.items()}
        row = {'number': number, 'values': values,
               'original': {str(i + 1): str(c.value) if c.value is not None else '' for i, c in enumerate(cells)},
               'error': ''}
        try:
            if any(c.data_type == 'f' for c in cells):
                raise ValueError('Formula cells are not accepted; upload values only.')
            values['code'] = employee_code(cells[mapping['code']])
            if not values['code']:
                raise ValueError('Missing employee code.')
        except ValueError as exc:
            row['error'] = str(exc)
        rows.append(row)
    return rows


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise ValueError('Invalid date; use an Excel date or YYYY-MM-DD.') from exc


def parse_time(value):
    if value is None or str(value).strip() == '':
        return None
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.tzinfo:
            raise ValueError('Timezone-aware Excel times are not supported.')
        return value
    if isinstance(value, (float, int)) and not isinstance(value, bool) and 0 <= value < 1:
        return (datetime(2000, 1, 1) + timedelta(seconds=round(value * 86400))).time()
    text = str(value).strip()
    if not re.fullmatch(r'\d{1,2}:\d{2}(:\d{2})?', text):
        raise ValueError('Invalid time: %s' % text)
    try:
        return datetime.strptime(text, '%H:%M:%S' if text.count(':') == 2 else '%H:%M').time()
    except ValueError as exc:
        raise ValueError('Invalid time: %s' % text) from exc


def to_utc(local, tz_name):
    zone = ZoneInfo(tz_name)
    candidates = set()
    for fold in (0, 1):
        candidate = local.replace(tzinfo=zone, fold=fold).astimezone(timezone.utc)
        if candidate.astimezone(zone).replace(tzinfo=None) == local:
            candidates.add(candidate.replace(tzinfo=None))
    if len(candidates) != 1:
        raise ValueError('Ambiguous or nonexistent local time (DST); correct the source time.')
    return candidates.pop()


def interval(values, tz_name, missing='strict', max_hours=16):
    day = parse_date(values.get('date'))
    punches = [values.get('check_in'), values.get('check_out')]
    if any(str(v).strip().lower() in LEAVE for v in punches):
        return {'status': 'skipped', 'message': 'Leave/absence text; no attendance created.', 'date': day}
    start, end = map(parse_time, punches)
    if not start or not end:
        if missing == 'skip':
            return {'status': 'skipped', 'message': 'Incomplete punches skipped.', 'date': day}
        if not start or missing != 'open':
            return {'status': 'error', 'message': 'Missing Check In.' if not start else 'Missing Check Out.', 'date': day}
    overnight = bool(end and end < start)
    check_in = to_utc(datetime.combine(day, start), tz_name)
    check_out = to_utc(datetime.combine(day + timedelta(days=int(overnight)), end), tz_name) if end else None
    duration = (check_out - check_in).total_seconds() / 3600 if check_out else 0
    if check_out and (duration <= 0 or duration > max_hours):
        raise ValueError('Shift duration must be positive and at most %s hours.' % max_hours)
    return {'date': day, 'check_in': check_in, 'check_out': check_out,
            'duration': duration, 'status': 'ready',
            'message': 'Overnight shift.' if overnight else ''}


def overlaps(a_start, a_end, b_start, b_end):
    return a_start < (b_end or datetime.max) and b_start < (a_end or datetime.max)
