import base64
from bisect import bisect_left, bisect_right
from collections import defaultdict
from datetime import datetime
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from odoo import api, fields, models, _
from odoo.exceptions import AccessError, UserError, ValidationError

from ..services.parser import interval, open_excel, overlaps, read_rows


_INTERNAL = object()


class Attendance(models.Model):
    _inherit = 'hr.attendance'

    excel_import_id = fields.Many2one('hr.attendance.excel.import', readonly=True, copy=False,
                                     ondelete='set null', groups='hr_attendance.group_hr_attendance_manager')

    @api.model_create_multi
    def create(self, vals_list):
        if any(v.get('excel_import_id') for v in vals_list) and self.env.context.get('_attendance_internal') is not _INTERNAL:
            raise AccessError(_('Import provenance can only be set by the importer.'))
        return super().create(vals_list)

    def write(self, vals):
        if 'excel_import_id' in vals and self.env.context.get('_attendance_internal') is not _INTERNAL:
            raise AccessError(_('Import provenance is immutable.'))
        return super().write(vals)


class ImportLock(models.Model):
    _name = 'hr.attendance.excel.import.lock'
    _description = 'Attendance Import Serialization Guard'

    revision = fields.Integer(default=0)

    def _load_records_write(self, values):
        # A forced init may replay the empty noupdate XML record. Odoo's
        # default loader calls write({}), which our mutation guard rejects.
        # This private loader hook accepts only the no-op declaration; it
        # neither resets the revision nor relaxes public write permissions.
        if not values:
            return
        return super()._load_records_write(values)

    def write(self, vals):
        if self.env.context.get('_attendance_internal') is not _INTERNAL:
            raise AccessError(_('Internal serialization guard.'))
        return super().write(vals)

    def unlink(self):
        raise AccessError(_('Internal serialization guard.'))


class AttendanceImport(models.Model):
    _name = 'hr.attendance.excel.import'
    _description = 'Attendance Excel Import History'
    _order = 'id desc'
    _rec_name = 'filename'

    filename = fields.Char(required=True, readonly=True)
    file_data = fields.Binary(required=True, attachment=False, readonly=True)
    company_id = fields.Many2one('res.company', required=True, readonly=True)
    user_id = fields.Many2one('res.users', default=lambda self: self.env.user, readonly=True)
    import_date = fields.Datetime(readonly=True)
    timezone = fields.Char(required=True, readonly=True)
    code_field = fields.Char(required=True, readonly=True)
    sheet_number = fields.Integer(default=1, readonly=True)
    duplicate_policy = fields.Selection([('skip', 'Skip Existing'), ('conflict', 'Report Conflicts'),
                                         ('update', 'Update Existing')], default='skip', required=True, readonly=True)
    missing_policy = fields.Selection([('strict', 'Strict'), ('open', 'Allow Open Attendance'),
                                       ('skip', 'Skip Incomplete')], default='strict', required=True, readonly=True)
    approve_manual = fields.Boolean(string='Approve updating records not created by this importer', readonly=True)
    max_hours = fields.Float(default=16, readonly=True)
    max_rows = fields.Integer(default=10000, readonly=True)
    max_mb = fields.Integer(default=10, readonly=True)
    state = fields.Selection([('preview', 'Preview'), ('done', 'Completed'), ('partial', 'Completed with errors')],
                             default='preview', readonly=True)
    line_ids = fields.One2many('hr.attendance.excel.import.line', 'import_id', readonly=True)
    total = fields.Integer(compute='_counts')
    imported = fields.Integer(compute='_counts')
    skipped = fields.Integer(compute='_counts')
    duplicates = fields.Integer(compute='_counts')
    failed = fields.Integer(compute='_counts')
    report_data = fields.Binary(attachment=False, readonly=True)
    report_name = fields.Char(default='attendance-errors.xlsx', readonly=True)

    @api.depends('line_ids.status')
    def _counts(self):
        for batch in self:
            batch.total = len(batch.line_ids)
            batch.imported = len(batch.line_ids.filtered(lambda r: r.status == 'imported'))
            batch.skipped = len(batch.line_ids.filtered(lambda r: r.status == 'skipped'))
            batch.duplicates = len(batch.line_ids.filtered(lambda r: r.status == 'duplicate'))
            batch.failed = len(batch.line_ids.filtered(lambda r: r.status == 'error'))

    def _authorize(self):
        if not self.env.user.has_group('hr_attendance.group_hr_attendance_manager'):
            raise AccessError(_('Only Attendance managers may import attendance.'))
        self.check_access('read')
        if any(b.company_id not in self.env.companies for b in self):
            raise AccessError(_('Select an allowed company.'))

    @api.model_create_multi
    def create(self, vals_list):
        if self.env.context.get('_attendance_internal') is not _INTERNAL:
            raise AccessError(_('Create imports using the upload wizard.'))
        records = super().create(vals_list)
        records._authorize()
        return records

    def write(self, vals):
        if self.env.context.get('_attendance_internal') is not _INTERNAL:
            self._authorize()
            for batch in self.sorted('id'):
                batch._lock()
            commands = vals.get('line_ids', [])
            if set(vals) != {'line_ids'} or any(b.state != 'preview' for b in self):
                raise AccessError(_('Import history cannot be edited directly.'))
            allowed_ids = set(self.line_ids.ids)
            if any(len(c) != 3 or c[0] != 1 or c[1] not in allowed_ids or set(c[2]) != {'approved'} for c in commands):
                raise AccessError(_('Only preview row approvals may be changed.'))
        return super().write(vals)

    def unlink(self):
        raise AccessError(_('Import audit history cannot be deleted.'))

    def _internal(self):
        return self.with_context(_attendance_internal=_INTERNAL)

    def _employee_model(self):
        employee = self.env['hr.employee'].with_context(active_test=False)
        available = employee.fields_get()
        definition = available.get(self.code_field, {})
        actual = employee._fields.get(self.code_field)
        if not actual or not actual.store or definition.get('type') != 'char' or actual.company_dependent:
            raise UserError(_('Choose a readable, stored text employee code field.'))
        employee.check_access('read')
        return employee

    def _parse(self):
        try:
            book = open_excel(base64.b64decode(self.with_context(bin_size=False).file_data, validate=True), self.max_mb)
            try:
                return read_rows(book, self.sheet_number, self.max_rows)
            finally:
                book.close()
        except ValueError as exc:
            raise UserError(str(exc)) from exc

    def _prepare(self):
        self.ensure_one()
        self._authorize()
        employee_model = self._employee_model()
        rows = self._parse()
        codes = {r['values'].get('code') for r in rows if not r['error']}
        employees = employee_model.search([(self.code_field, 'in', list(codes)), ('company_id', '=', self.company_id.id)])
        matched = defaultdict(list)
        for employee in employees:
            matched[employee[self.code_field]].append(employee)
        prepared = []
        for row in rows:
            values = row['values']
            result = {'import_id': self.id, 'row_number': row['number'],
                      'code': str(values.get('code') or ''), 'shift': str(values.get('shift') or ''),
                      'original': json.dumps(row['original'], ensure_ascii=False),
                      'status': 'error', 'message': row['error']}
            candidates = matched.get(values.get('code'), [])
            if not row['error']:
                if len(candidates) != 1:
                    result['message'] = 'Missing Employee.' if not candidates else 'Ambiguous Employee.'
                else:
                    result['employee_id'] = candidates[0].id
                    result['employee_name'] = candidates[0].name
                    try:
                        result.update(interval(values, self.timezone, self.missing_policy, self.max_hours))
                    except ValueError as exc:
                        result['message'] = str(exc)
            prepared.append(result)
        self._classify(prepared)
        return prepared

    def _classify(self, rows):
        ready = [r for r in rows if r['status'] == 'ready']
        if not ready:
            return
        attendance = self.env['hr.attendance']
        attendance.check_access('read')
        lower = min(r['check_in'] for r in ready)
        upper = max((r.get('check_out') or datetime.max) for r in ready)
        records = attendance.search([
            ('employee_id', 'in', list({r['employee_id'] for r in ready})),
            ('check_in', '<=', upper), '|', ('check_out', '=', False), ('check_out', '>=', lower),
        ])
        existing = defaultdict(list)
        for record in records:
            existing[record.employee_id.id].append(record)
        indexes = {}
        for employee_id, items in existing.items():
            items.sort(key=lambda r: r.check_in)
            starts, max_ends, by_start = [], [], defaultdict(list)
            for record in items:
                starts.append(record.check_in)
                max_ends.append(max(max_ends[-1] if max_ends else datetime.min, record.check_out or datetime.max))
                by_start[record.check_in].append(record)
            indexes[employee_id] = (starts, max_ends, by_start)
        accepted = defaultdict(list)
        for row in sorted(ready, key=lambda r: (r['check_in'], r['row_number'])):
            employee = row['employee_id']
            start, end = row['check_in'], row.get('check_out') or False
            prior = accepted[employee]
            if prior and start == prior[-1]['check_in'] and end == (prior[-1].get('check_out') or False):
                row.update(status='duplicate', message='Duplicate row within this file.')
                continue
            if prior and overlaps(start, end, prior[-1]['check_in'], prior[-1].get('check_out')):
                row.update(status='error', message='Overlapping attendance within this file.')
                continue
            starts, max_ends, by_start = indexes.get(employee, ([], [], {}))
            same = by_start.get(start, [])
            target = same[0] if len(same) == 1 else None
            if target and target.check_out == end:
                status = 'error' if self.duplicate_policy == 'conflict' else 'duplicate'
                row.update(status=status, message='Identical attendance already exists.')
                continue
            if same:
                if not target or self.duplicate_policy != 'update':
                    row.update(status='error', message='Existing Check In has different Check Out.')
                    continue
                if not target.excel_import_id and not self.approve_manual:
                    row.update(status='error', message='Explicit approval required to update attendance not created by this importer.')
                    continue
                # Never erase a recorded punch to reopen a completed session.
                if target.check_out and not end:
                    row.update(status='error', message='Cannot clear an existing Check Out.')
                    continue
                row['target_id'] = target.id
                row['previous_check_out'] = target.check_out
            candidates = existing[employee][bisect_right(max_ends, start):bisect_left(starts, end or datetime.max)]
            if any(r != target and overlaps(start, end, r.check_in, r.check_out) for r in candidates):
                row.update(status='error', message='Overlapping existing attendance.')
                continue
            if target:
                row['message'] = 'Update attendance #%s. Previous Check Out: %s. %s' % (target.id, target.check_out, row['message'])
            prior.append(row)

    def action_validate(self):
        self.ensure_one()
        self._authorize()
        self._lock()
        if self.state != 'preview':
            raise UserError(_('This import is already complete.'))
        prepared = self._prepare()
        self.line_ids._internal().unlink()
        self.env['hr.attendance.excel.import.line'].with_context(_attendance_internal=_INTERNAL).create(prepared)
        return self._open()

    def _open(self):
        return {'type': 'ir.actions.act_window', 'res_model': self._name, 'res_id': self.id,
                'view_mode': 'form', 'target': 'current'}

    def action_import(self):
        self.ensure_one()
        self._authorize()
        self._lock()
        if self.state != 'preview' or not self.line_ids:
            raise UserError(_('Validate and preview before confirming an import.'))
        self.env['hr.attendance'].check_access('create')
        return self._execute_import()

    def _lock(self):
        self.ensure_one()
        # Updating one guard serializes all importer transactions, including different files.
        # A stale repeatable-read snapshot produces a serialization failure, retried by Odoo.
        guard = self.env.ref('hr_attendance_excel_import.serialization_guard')
        self.env.cr.execute('SELECT id FROM hr_attendance_excel_import_lock WHERE id = %s FOR UPDATE', [guard.id])
        guard.invalidate_recordset()
        guard.with_context(_attendance_internal=_INTERNAL).write({'revision': guard.revision + 1})
        guard.flush_recordset()
        self.env.cr.execute('SELECT id FROM hr_attendance_excel_import WHERE id = %s FOR UPDATE', [self.id])
        self.invalidate_recordset()

    def _execute_import(self):
        # Reparse trusted source and recheck current permissions/conflicts, never trust editable preview values.
        current = {r['row_number']: r for r in self._prepare()}
        for line in self.line_ids.sorted(lambda r: (r.check_in or datetime.max, r.row_number)):
            fresh = current[line.row_number]
            if not line.approved:
                line._internal().write({'status': 'skipped', 'message': 'Deselected by importing manager.'})
                continue
            if line.status != 'ready' or fresh['status'] != 'ready':
                line._internal().write({'status': fresh['status'] if fresh['status'] != 'ready' else 'error',
                                       'message': fresh['message'] or 'Revalidate this row before importing.'})
                continue
            fingerprint = (line.employee_id.id, line.check_in, line.check_out, line.target_id.id, line.previous_check_out)
            if fingerprint != (fresh['employee_id'], fresh['check_in'], fresh.get('check_out') or False, fresh.get('target_id', False), fresh.get('previous_check_out', False)):
                line._internal().write({'status': 'error', 'message': 'Matching attendance changed since preview; create a new preview.'})
                continue
            try:
                with self.env.cr.savepoint():
                    vals = {'employee_id': fresh['employee_id'], 'check_in': fresh['check_in'],
                            'check_out': fresh.get('check_out') or False}
                    if fresh.get('target_id'):
                        record = self.env['hr.attendance'].browse(fresh['target_id'])
                        before = str(record.check_out)
                        record.write({'check_out': vals['check_out']})
                        message = 'Updated #%s; previous Check Out: %s.' % (record.id, before)
                    else:
                        record = self.env['hr.attendance'].with_context(_attendance_internal=_INTERNAL).create(dict(vals, excel_import_id=self.id))
                        message = 'Created #%s.' % record.id
                    line._internal().write({'status': 'imported', 'attendance_id': record.id, 'message': message})
            except ValidationError as exc:
                line._internal().write({'status': 'error', 'message': str(exc)})
        self._internal().write({'state': 'partial' if any(r.status == 'error' for r in self.line_ids) else 'done',
                                'user_id': self.env.user.id, 'import_date': fields.Datetime.now()})
        return self._open()

    def action_error_report(self):
        self.ensure_one()
        self._authorize()
        # Read the selected sheet from the retained upload, also for older
        # imports whose audit rows only store numbered original columns.
        try:
            source = open_excel(
                base64.b64decode(self.with_context(bin_size=False).file_data, validate=True), self.max_mb)
            try:
                if not 1 <= self.sheet_number <= len(source.worksheets):
                    raise ValueError('Select a valid worksheet number.')
                headers = list(next(source.worksheets[self.sheet_number - 1].iter_rows(
                    min_row=1, max_row=1, values_only=True), ()))
            finally:
                source.close()
        except ValueError as exc:
            raise UserError(str(exc)) from exc
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet('Errors')
        lines = self.line_ids.filtered(lambda r: r.status != 'imported')
        originals = {line.id: json.loads(line.original) for line in lines}
        columns = max(len(headers), max((len(row) for row in originals.values()), default=0))
        headers.extend([None] * (columns - len(headers)))
        # Preserve spelling, whitespace, duplicate/blank headings and order.
        # Explicit string cells prevent a heading beginning '=' becoming a formula.
        header_cells = []
        for value in headers:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = 's'
            header_cells.append(cell)
        sheet.append(['Row', 'Employee code', 'Status', 'Reason'] + header_cells)
        def safe(value):
            text = str(value or '')
            return "'" + text if text.startswith(('=', '+', '-', '@')) else text
        for line in lines:
            sheet.append([line.row_number, safe(line.code), line.status, safe(line.message)] +
                         [safe(originals[line.id].get(str(i))) for i in range(1, columns + 1)])
        output = BytesIO()
        workbook.save(output)
        self._internal().write({'report_data': base64.b64encode(output.getvalue())})
        return {'type': 'ir.actions.act_url', 'url': '/web/content/%s/%s/report_data/%s?download=true' % (
            self._name, self.id, self.report_name), 'target': 'self'}


class AttendanceImportLine(models.Model):
    _name = 'hr.attendance.excel.import.line'
    _description = 'Attendance Import Row Result'
    _order = 'row_number, id'

    import_id = fields.Many2one('hr.attendance.excel.import', required=True, ondelete='cascade', index=True)
    company_id = fields.Many2one(related='import_id.company_id', store=True, index=True)
    row_number = fields.Integer()
    code = fields.Char()
    employee_id = fields.Many2one('hr.employee', ondelete='set null')
    employee_name = fields.Char()
    date = fields.Date()
    shift = fields.Char()
    check_in = fields.Datetime()
    check_out = fields.Datetime()
    duration = fields.Float()
    status = fields.Selection([('ready', 'Ready'), ('duplicate', 'Duplicate'), ('error', 'Error'),
                               ('skipped', 'Skipped'), ('imported', 'Imported')])
    message = fields.Text()
    original = fields.Text()
    approved = fields.Boolean(default=True)
    target_id = fields.Many2one('hr.attendance', ondelete='set null')
    previous_check_out = fields.Datetime()
    attendance_id = fields.Many2one('hr.attendance', ondelete='set null')

    def _internal(self):
        return self.with_context(_attendance_internal=_INTERNAL)

    @api.model_create_multi
    def create(self, vals_list):
        if self.env.context.get('_attendance_internal') is not _INTERNAL:
            raise AccessError(_('Preview rows are generated by validation.'))
        return super().create(vals_list)

    def write(self, vals):
        if self.env.context.get('_attendance_internal') is not _INTERNAL:
            batches = self.mapped('import_id')
            batches._authorize()
            for batch in batches.sorted('id'):
                batch._lock()
            if set(vals) != {'approved'} or any(r.import_id.state != 'preview' for r in self):
                raise AccessError(_('Only row approval can be edited before import.'))
        return super().write(vals)

    def unlink(self):
        if self.env.context.get('_attendance_internal') is not _INTERNAL:
            raise AccessError(_('Audit rows cannot be deleted.'))
        return super().unlink()
