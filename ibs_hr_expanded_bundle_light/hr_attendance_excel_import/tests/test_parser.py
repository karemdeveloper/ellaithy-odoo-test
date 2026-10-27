"""Run without Odoo: python3 tests/test_parser.py."""
from datetime import date, datetime, time
import importlib.util
from io import BytesIO
from pathlib import Path
import unittest

from openpyxl import Workbook

spec = importlib.util.spec_from_file_location('attendance_parser', Path(__file__).resolve().parents[1] / 'services/parser.py')
parser = importlib.util.module_from_spec(spec)
spec.loader.exec_module(parser)


def workbook_bytes(rows, headers=None):
    book = Workbook()
    sheet = book.active
    sheet.append(headers or ['كود الموظف', 'التاريخ', 'الحضور', 'الانصراف'])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class ParserTests(unittest.TestCase):
    def values(self, start='08:00', end='16:00', day='2026-09-02'):
        return {'date': day, 'check_in': start, 'check_out': end}

    def test_valid_cairo_conversion(self):
        row = parser.interval(self.values(), 'Africa/Cairo')
        self.assertEqual(row['check_in'], datetime(2026, 9, 2, 5))
        self.assertEqual(row['duration'], 8)

    def test_overnight(self):
        row = parser.interval(self.values('13:48:22', '00:05:44'), 'Africa/Cairo')
        self.assertEqual(row['check_out'], datetime(2026, 9, 2, 21, 5, 44))
        self.assertEqual(row['message'], 'Overnight shift.')

    def test_missing_check_in(self):
        self.assertEqual(parser.interval(self.values(None), 'UTC', 'open')['status'], 'error')

    def test_missing_check_out(self):
        self.assertEqual(parser.interval(self.values(end=None), 'UTC')['status'], 'error')

    def test_open(self):
        self.assertIsNone(parser.interval(self.values(end=None), 'UTC', 'open')['check_out'])

    def test_skip_incomplete(self):
        self.assertEqual(parser.interval(self.values(end=None), 'UTC', 'skip')['status'], 'skipped')

    def test_leave(self):
        for text in ('إجازة', 'اجازة', 'Leave'):
            self.assertEqual(parser.interval(self.values(text, text), 'UTC')['status'], 'skipped')

    def test_invalid_time(self):
        for value in ('bad', '25:00', '12:61', -0.1, 1.5, True):
            with self.subTest(value=value), self.assertRaises(ValueError):
                parser.parse_time(value)

    def test_time_types(self):
        for value in (time(12), datetime(2026, 9, 2, 12), ' 12:00 ', '12:00:00', 0.5):
            self.assertEqual(parser.parse_time(value), time(12))

    def test_date_types(self):
        for value in (date(2026, 9, 2), datetime(2026, 9, 2), '2026-09-02'):
            self.assertEqual(parser.parse_date(value), date(2026, 9, 2))

    def test_invalid_date(self):
        with self.assertRaises(ValueError):
            parser.parse_date('02/09/2026')

    def test_dst_ambiguous(self):
        with self.assertRaises(ValueError):
            parser.to_utc(datetime(2026, 11, 1, 1, 30), 'America/New_York')

    def test_dst_nonexistent(self):
        with self.assertRaises(ValueError):
            parser.to_utc(datetime(2026, 3, 8, 2, 30), 'America/New_York')

    def test_duration_limits(self):
        for start, end in [('08:00', '08:00'), ('08:00', '07:00')]:
            with self.assertRaises(ValueError):
                parser.interval(self.values(start, end), 'UTC')

    def test_headers_and_empty_rows(self):
        data = workbook_bytes([['EMP-892', '2026-09-02', '08:00', '16:00'], [None] * 4])
        book = parser.open_excel(data)
        self.addCleanup(book.close)
        self.assertEqual(len(parser.read_rows(book)), 1)

    def test_english_headers(self):
        book = parser.open_excel(workbook_bytes([['EMP-892', '2026-09-02', '08:00', '16:00']],
                                                ['Employee Code', 'Date', 'Check In', 'Check Out']))
        self.addCleanup(book.close)
        self.assertEqual(parser.read_rows(book)[0]['values']['code'], 'EMP-892')

    def test_formula_rejected(self):
        book = parser.open_excel(workbook_bytes([['=1+1', '2026-09-02', '08:00', '16:00']]))
        self.addCleanup(book.close)
        self.assertIn('Formula', parser.read_rows(book)[0]['error'])

    def test_leading_zeros(self):
        book = Workbook()
        self.addCleanup(book.close)
        cell = book.active['A1']
        cell.value = 12
        cell.number_format = '00000'
        self.assertEqual(parser.employee_code(cell), '00012')
        cell.value = ' 0012 '
        self.assertEqual(parser.employee_code(cell), '0012')

    def test_missing_headers(self):
        book = parser.open_excel(workbook_bytes([], ['Code']))
        self.addCleanup(book.close)
        with self.assertRaises(ValueError):
            parser.read_rows(book)

    def test_malformed(self):
        with self.assertRaises(ValueError):
            parser.open_excel(b'not an Excel file')

    def test_oversized(self):
        with self.assertRaises(ValueError):
            parser.open_excel(b'x' * (1024 * 1024 + 1), max_mb=1)

    def test_row_limit(self):
        book = parser.open_excel(workbook_bytes([['A', '2026-09-02', '08:00', '16:00']] * 2))
        self.addCleanup(book.close)
        with self.assertRaises(ValueError):
            parser.read_rows(book, max_rows=1)

    def test_multiple_sheets(self):
        book = Workbook()
        self.addCleanup(book.close)
        book.create_sheet('Second').append(['Code', 'Date', 'Check In', 'Check Out'])
        self.assertEqual(parser.read_rows(book, 2), [])
        with self.assertRaises(ValueError):
            parser.read_rows(book, 3)

    def test_overlap_boundaries(self):
        a, b, c = [datetime(2026, 1, 1, h) for h in (8, 12, 16)]
        self.assertFalse(parser.overlaps(a, b, b, c))
        self.assertTrue(parser.overlaps(a, None, b, c))
        self.assertTrue(parser.overlaps(a, c, b, c))


if __name__ == '__main__':
    unittest.main(verbosity=2)
